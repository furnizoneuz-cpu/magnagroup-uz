#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""ASTATKA.xlsx (13.07.2026) — real showroom qoldig'i:
har varaqdan qator ma'lumotlari + shu qatorga bog'langan rasmlarni chiqaradi.
Natija: data/astatka.json + public/products/real/*.jpg (rasmlar artikul nomi bilan)."""
import os, json, re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Katalog", "Rasmli malumotlar", "ASTATKA.xlsx13.07.2026.xlsx")
IMGDIR = os.path.join(ROOT, "public", "products", "real")
OUT = os.path.join(ROOT, "data", "astatka.json")
os.makedirs(IMGDIR, exist_ok=True)

def slug(s):
    s = str(s).strip()
    s = re.sub(r"[^\w\-\.]+", "_", s, flags=re.UNICODE)
    return s[:60] or "item"

wb = openpyxl.load_workbook(XLSX, data_only=True)

def cells(ws):
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                grid[(c.row, c.column)] = c.value
    return grid

def images_by_row(ws):
    """anchor row (1-based) -> list of image objects"""
    out = {}
    for img in getattr(ws, "_images", []):
        try:
            r = img.anchor._from.row + 1
        except Exception:
            continue
        out.setdefault(r, []).append(img)
    return out

def save_img(img, name):
    ext = (img.format or "png").lower()
    if ext == "jpeg": ext = "jpg"
    path = os.path.join(IMGDIR, f"{name}.{ext}")
    data = img._data()
    with open(path, "wb") as f:
        f.write(data)
    return f"/products/real/{name}.{ext}"

result = {}

# ---- Sheet: Остатка До Август Бобомурод (kreslolar, rang bilan) ----
ws = wb["Остатка До Август Бобомурод"]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    name = g.get((r, 2))
    if not name: continue
    stock = g.get((r, 4))
    img = None
    if r in ib:
        img = save_img(ib[r][0], "ost_" + slug(name))
    rows.append({"name_ru": str(name).strip(), "stock": stock if isinstance(stock,(int,float)) else 0, "image": img})
# rangi boshqa juft qatorlar bitta rasmga ega — bo'sh qolganiga yaqinidagi (bazasi mos) rasmni beramiz
base = lambda s: re.split(r"\s+", s)[0]
for i, row in enumerate(rows):
    if row["image"]: continue
    for j in (i - 1, i + 1):
        if 0 <= j < len(rows) and rows[j]["image"] and base(rows[j]["name_ru"]) == base(row["name_ru"]):
            row["image"] = rows[j]["image"]; break
result["ostatka_kreslo"] = rows

# ---- Sheet: 1-Партия мебель+Кресло (to'liq: model/desc/size/color/stock/price USD) ----
ws = wb["1-Партия мебель+Кресло"]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    model = g.get((r, 2))
    if not model or str(model).strip() == "Модель": continue
    img = None
    if r in ib:
        img = save_img(ib[r][0], "p1_" + slug(model))
    rows.append({
        "article": str(model).strip(),
        "desc_ru": str(g.get((r, 3)) or "").strip(),
        "dimensions": str(g.get((r, 4)) or "").strip(),
        "color_ru": str(g.get((r, 5)) or "").strip(),
        "stock": g.get((r, 7)) if isinstance(g.get((r, 7)), (int, float)) else 0,
        "priceUSD": g.get((r, 8)) if isinstance(g.get((r, 8)), (int, float)) else None,
        "image": img,
    })
result["partiya1"] = rows

# ---- Sheet: 2-Партия Мебель+кресло (model/photo/stock) ----
ws = wb["2-Партия Мебель+кресло"]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    model = g.get((r, 2))
    if not model: continue
    img = None
    if r in ib:
        img = save_img(ib[r][0], "p2m_" + slug(model) + f"_{r}")
    rows.append({"article": str(model).strip(),
                 "stock": g.get((r, 4)) if isinstance(g.get((r, 4)), (int, float)) else 0,
                 "image": img})
result["partiya2_mebel"] = rows

# ---- Sheet: 2-Партия Кресло (G-kodlar: № | Фото | Модель | stock) ----
ws = wb["2-Партия Кресло"]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    model = g.get((r, 3))
    if not model: continue
    img = None
    if r in ib:
        img = save_img(ib[r][0], "g_" + slug(model) + f"_{r}")
    rows.append({"article": str(model).strip(),
                 "stock": g.get((r, 4)) if isinstance(g.get((r, 4)), (int, float)) else 0,
                 "image": img})
result["partiya2_kreslo"] = rows

# ---- Sheet: 2-партия диван (SF divanlar; artikul A ustunda, davom qatorlari otasiga qo'shiladi) ----
ws = wb["2-партия диван"]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    art = g.get((r, 1))
    desc = g.get((r, 3))
    if art and str(art).strip() in ("Артикул", "№"): continue
    if not art and not desc: continue
    img = None
    if r in ib:
        name = slug(art) if art else f"row{r}"
        img = save_img(ib[r][0], "sf_" + name)
    if art:
        rows.append({"article": str(art).strip(),
                     "desc_ru": str(desc or "").strip(),
                     "material_ru": str(g.get((r, 4)) or "").strip(),
                     "stock": g.get((r, 6)) if isinstance(g.get((r, 6)), (int, float)) else None,
                     "image": img})
    elif rows:
        # davom qatori — otasidagi tavsifga qo'shimcha (3 kishilik varianti va h.k.)
        parent = rows[-1]
        if desc: parent["desc_ru"] += "\n+ " + str(desc).strip()
        if img and not parent["image"]: parent["image"] = img
result["partiya2_divan"] = rows

# ---- Sheet: 3-partiya divan 2026 ----
ws = wb["3-partiya divan 2026 "]
g = cells(ws); ib = images_by_row(ws)
rows = []
for r in range(2, ws.max_row + 1):
    art = g.get((r, 1))
    stock = g.get((r, 4))
    img = None
    if r in ib:
        name = slug(art) if art else f"row{r}"
        img = save_img(ib[r][0], "sf3_" + name)
    if not art and not img and stock is None: continue
    rows.append({"article": str(art).strip() if art else None,
                 "stock": stock if isinstance(stock, (int, float)) else None,
                 "image": img})
result["partiya3_divan"] = rows

json.dump(result, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
n_img = len([f for f in os.listdir(IMGDIR)])
print("Varaqlar:", {k: len(v) for k, v in result.items()})
print("Rasmlar saqlandi:", n_img, "->", IMGDIR)
print("JSON:", OUT)
